"""Slide-deck grader (T07): run the model's python-pptx code, then inspect the
generated deck.pptx — slide count, required title substrings, and a chart on the
required slide. Structural only; content quality is left to a rubric if desired."""
from __future__ import annotations

import pathlib
import subprocess
import sys
import tempfile

from .base import GradeContext, GradeResult, extract_code


def grade(task: dict, response_text: str, ctx: GradeContext) -> GradeResult:
    args = task.get("grader_args", {})
    n_slides = int(args.get("n_slides", 6))
    required_titles = [t.lower() for t in args.get("required_titles", [])]
    chart_slide = int(args.get("chart_slide", 3))          # 1-indexed
    code = extract_code(response_text, "python")
    if not code:
        return GradeResult(False, 0.0, "no python code found")

    with tempfile.TemporaryDirectory() as td:
        d = pathlib.Path(td)
        (d / "gen.py").write_text(code)
        proc = subprocess.run([sys.executable, "gen.py"], cwd=d,
                              capture_output=True, text=True, timeout=90)
        deck = d / "deck.pptx"
        if not deck.exists():
            # accept any .pptx the code produced
            found = list(d.glob("*.pptx"))
            if not found:
                return GradeResult(False, 0.0,
                                   f"code ran but produced no .pptx (rc={proc.returncode}): "
                                   f"{proc.stderr[-150:]}")
            deck = found[0]

        from pptx import Presentation  # noqa: PLC0415
        from pptx.util import Emu  # noqa: PLC0415,F401
        prs = Presentation(str(deck))
        slides = prs.slides
        count = len(slides)

        titles_text = []
        for s in slides:
            for shp in s.shapes:
                if shp.has_text_frame:
                    titles_text.append(shp.text_frame.text.lower())
        alltext = " ".join(titles_text)
        titles_found = [t for t in required_titles if t in alltext]

        has_chart = False
        if 1 <= chart_slide <= count:
            for shp in list(slides)[chart_slide - 1].shapes:
                if getattr(shp, "has_chart", False):
                    has_chart = True
                    break

    checks = {
        "slide_count": count == n_slides,
        "titles": len(titles_found) == len(required_titles),
        "chart": has_chart,
    }
    score = sum(checks.values()) / len(checks)
    passed = all(checks.values())
    return GradeResult(passed, score,
                       f"slides={count}/{n_slides} titles={len(titles_found)}/{len(required_titles)} "
                       f"chart_on_{chart_slide}={has_chart}")
